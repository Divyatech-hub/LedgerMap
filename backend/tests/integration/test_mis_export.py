from io import BytesIO
from pathlib import Path

import openpyxl
from fastapi.testclient import TestClient

FIXTURE = Path(__file__).parents[3] / "prototype" / "TB.xlsx"


def _create_client(db_client: TestClient, name: str) -> int:
    response = db_client.post("/clients", json={"name": name})
    assert response.status_code == 201, response.text
    return response.json()["id"]


def _upload_run(db_client: TestClient, client_id: int, period: str) -> dict:
    with FIXTURE.open("rb") as fixture_file:
        response = db_client.post(
            f"/runs/clients/{client_id}",
            data={"period": period},
            files={"file": ("TB.xlsx", fixture_file, "application/octet-stream")},
        )
    assert response.status_code == 201, response.text
    return response.json()


def test_export_for_client_with_no_runs_returns_404(db_client: TestClient) -> None:
    client_id = _create_client(db_client, "No Runs Co")

    response = db_client.get(f"/exports/clients/{client_id}/mis")

    assert response.status_code == 404


def test_export_returns_a_workbook_with_summary_formulas_and_line_items(
    db_client: TestClient,
) -> None:
    client_id = _create_client(db_client, "Export Verify Co")
    _upload_run(db_client, client_id, period="2024-01")

    response = db_client.get(f"/exports/clients/{client_id}/mis")

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "Export Verify Co" in response.headers["content-disposition"]

    workbook = openpyxl.load_workbook(BytesIO(response.content), data_only=False)
    assert "MIS Summary" in workbook.sheetnames
    assert "Line Items 2024-01" in workbook.sheetnames

    summary = workbook["MIS Summary"]
    assert [c.value for c in summary[1]][:3] == ["Code", "Description", "2024-01"]

    # Find the row for code 2068 and confirm its cell is a live SUMIF formula
    # against the Line Items sheet, not a precomputed value.
    code_column = [cell.value for cell in summary["A"]]
    row_index = code_column.index("2068") + 1
    formula = summary.cell(row=row_index, column=3).value
    assert isinstance(formula, str)
    assert formula.startswith("=SUMIF('Line Items 2024-01'!")
    assert f"$A{row_index}" in formula

    line_items = workbook["Line Items 2024-01"]
    headers = [c.value for c in line_items[1]]
    assert headers == [
        "Account",
        "Ancestors",
        "Amount",
        "Code",
        "Method",
        "MIS formula",
    ]

    # At least one resolved line item should carry its own SUMIF formula
    # text, so an accountant can see exactly what feeds the summary total.
    formula_column = [cell.value for cell in line_items["F"]][1:]
    assert any(
        isinstance(value, str) and value.startswith("=SUMIF(")
        for value in formula_column
    )


def test_export_reflects_a_correction(db_client: TestClient) -> None:
    client_id = _create_client(db_client, "Correction Export Co")
    run = _upload_run(db_client, client_id, period="2024-01")
    review_item = next(item for item in run["line_items"] if item["status"] == "review")

    correction = db_client.post(
        f"/runs/{run['id']}/line-items/{review_item['id']}/corrections",
        json={"resulting_code": "9001"},
    )
    assert correction.status_code == 201, correction.text

    response = db_client.get(f"/exports/clients/{client_id}/mis")
    assert response.status_code == 200

    workbook = openpyxl.load_workbook(BytesIO(response.content), data_only=False)
    summary = workbook["MIS Summary"]
    code_column = [cell.value for cell in summary["A"]]
    assert "9001" in code_column
