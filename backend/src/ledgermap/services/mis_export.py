"""Write the Detailed MIS preview out as a real, downloadable .xlsx.

Per the README's product surface: "writes the real .xlsx ... matching the
client's existing template/formulas exactly." Reverse-engineering an
arbitrary client's own template turned out to be a much larger, riskier
feature (see the design discussion this followed) — this instead writes a
workbook we fully control: one "MIS Summary" sheet whose cells are real,
Excel-native `SUMIF` formulas (not precomputed values, so Excel visibly
recalculates them), sourced from one "Line Items <period>" sheet per period
that lists every original account with the exact formula that sums it (and
its siblings under the same code) into the summary — so an accountant can
click any summary cell, jump to its source rows, and check the arithmetic
themselves, the same trust an in-house template would give them.
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from ledgermap.db.models import Run
from ledgermap.services.mis_report import MisReport, build_mis_report

_HEADER_FONT = Font(bold=True)
_REVIEW_FILL_ARGB = "FFF4E0C4"
_CONFIDENT_FILL_ARGB = "FFDCEEE1"

# Line Items sheet column layout, used both when writing rows and when
# building the SUMIF formulas that reference them.
_LI_ACCOUNT_COL = 1
_LI_ANCESTORS_COL = 2
_LI_AMOUNT_COL = 3
_LI_CODE_COL = 4
_LI_METHOD_COL = 5
_LI_FORMULA_COL = 6


def _line_items_sheet_name(period: str) -> str:
    # Excel sheet names cap at 31 chars and can't contain []:*?/\
    safe_period = "".join(c for c in period if c not in "[]:*?/\\")
    return f"Line Items {safe_period}"[:31]


def _write_line_items_sheet(workbook: Workbook, run: Run) -> None:
    sheet = workbook.create_sheet(_line_items_sheet_name(run.period))
    headers = ["Account", "Ancestors", "Amount", "Code", "Method", "MIS formula"]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = _HEADER_FONT

    code_col_letter = get_column_letter(_LI_CODE_COL)
    amount_col_letter = get_column_letter(_LI_AMOUNT_COL)
    sheet_ref = f"'{sheet.title}'"

    for item in run.line_items:
        row_index = sheet.max_row + 1
        sheet.append(
            [
                item.raw_name,
                " > ".join(item.ancestors) if item.ancestors else None,
                float(item.amount),
                item.matched_code,
                item.method,
                None,
            ]
        )
        if item.matched_code is not None:
            formula = (
                f"=SUMIF({sheet_ref}!${code_col_letter}:${code_col_letter},"
                f"{code_col_letter}{row_index},"
                f"{sheet_ref}!${amount_col_letter}:${amount_col_letter})"
            )
            sheet.cell(row=row_index, column=_LI_FORMULA_COL, value=formula)

    for column_cells in sheet.columns:
        length = max(
            (len(str(c.value)) for c in column_cells if c.value is not None), default=8
        )
        sheet.column_dimensions[column_cells[0].column_letter].width = min(
            length + 2, 60
        )


def _write_summary_sheet(workbook: Workbook, report: MisReport) -> Worksheet:
    sheet = workbook.create_sheet("MIS Summary", 0)
    headers = ["Code", "Description", *report.periods]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = _HEADER_FONT

    code_col_letter = get_column_letter(_LI_CODE_COL)
    amount_col_letter = get_column_letter(_LI_AMOUNT_COL)

    for row in report.rows:
        row_index = sheet.max_row + 1
        sheet.cell(row=row_index, column=1, value=row.code)
        sheet.cell(row=row_index, column=2, value=row.description)

        for period_index, period in enumerate(report.periods):
            column = 3 + period_index
            cell = sheet.cell(row=row_index, column=column)
            source_sheet_ref = f"'{_line_items_sheet_name(period)}'"
            fill_argb = None
            mis_cell = row.cells.get(period)
            if mis_cell is not None:
                cell.value = (
                    f"=SUMIF({source_sheet_ref}!${code_col_letter}:${code_col_letter},"
                    f"$A{row_index},"
                    f"{source_sheet_ref}!${amount_col_letter}:${amount_col_letter})"
                )
                fill_argb = (
                    _REVIEW_FILL_ARGB
                    if mis_cell.method == "review"
                    else _CONFIDENT_FILL_ARGB
                )
            if fill_argb:
                cell.fill = PatternFill(
                    start_color=fill_argb, end_color=fill_argb, fill_type="solid"
                )

    for column_cells in sheet.columns:
        length = max(
            (len(str(c.value)) for c in column_cells if c.value is not None), default=8
        )
        sheet.column_dimensions[column_cells[0].column_letter].width = min(
            length + 2, 40
        )

    return sheet


def build_mis_workbook(
    runs: list[Run], taxonomy: dict[str, str] | None = None
) -> bytes:
    """Build the exportable .xlsx for a client's full run history.

    `runs` must be ordered oldest-to-newest with line items loaded (see
    `list_runs_for_client_with_line_items`), matching what the MIS preview
    endpoint uses, so the export always matches what was last shown on
    screen.
    """
    report = build_mis_report(runs, taxonomy=taxonomy)

    workbook = Workbook()
    # Remove the default blank sheet; summary and line-item sheets are added
    # explicitly below.
    workbook.remove(workbook.active)

    _write_summary_sheet(workbook, report)
    for run in runs:
        _write_line_items_sheet(workbook, run)

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
