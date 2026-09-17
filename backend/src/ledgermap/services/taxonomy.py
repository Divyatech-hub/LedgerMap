from pathlib import Path
from io import BytesIO
from typing import BinaryIO

import openpyxl


def extract_taxonomy(
    source: str | Path | bytes | BinaryIO,
    *,
    sheet_name: str = "Detailed MIS",
) -> dict[str, str]:
    workbook = openpyxl.load_workbook(
        BytesIO(source) if isinstance(source, bytes) else source,
        read_only=True,
        data_only=True,
    )
    worksheet = workbook[sheet_name] if sheet_name in workbook.sheetnames else _find_mis_sheet(workbook)
    taxonomy: dict[str, str] = {}
    for code, description in worksheet.iter_rows(min_row=1, max_col=2, values_only=True):
        if code in (None, "") or description in (None, ""):
            continue
        code_text = str(code).strip()
        if code_text.isdigit():
            taxonomy.setdefault(code_text, str(description).strip())
    return taxonomy


def _find_mis_sheet(workbook: openpyxl.Workbook):
    for worksheet in workbook.worksheets:
        title = worksheet.title.casefold()
        if "detailed mis" in title or title == "mis":
            return worksheet
    raise KeyError("No Detailed MIS sheet found")