from collections.abc import Iterable
from io import BytesIO
from pathlib import Path
from typing import BinaryIO

import openpyxl

from ledgermap.domain.models import TrialBalanceRow
from ledgermap.services.tb_parser import amount_from_value


def read_trial_balance(
    source: str | Path | bytes | BinaryIO,
    *,
    sheet_name: str | None = None,
) -> list[TrialBalanceRow]:
    workbook = openpyxl.load_workbook(_as_workbook_source(source), read_only=True, data_only=True)
    worksheet = workbook[sheet_name] if sheet_name else _find_trial_balance_sheet(workbook)
    rows: list[TrialBalanceRow] = []
    for row_number, values in enumerate(worksheet.iter_rows(values_only=False), start=1):
        name_cell = values[0] if values else None
        name = str(name_cell.value).strip() if name_cell and name_cell.value else ""
        if not name:
            continue
        amount_cell = values[4] if len(values) > 4 else None
        amount = amount_from_value(amount_cell.value if amount_cell else None)
        rows.append(
            TrialBalanceRow(
                name=name,
                amount=amount,
                row_number=row_number,
                depth=_cell_depth(name_cell),
                is_bold=bool(name_cell.font.bold) if name_cell else False,
            )
        )
    return rows


def read_mapping_memory(
    source: str | Path | bytes | BinaryIO,
    *,
    sheet_name: str = "IFRS mapping sheet",
) -> dict[str, str | None]:
    workbook = openpyxl.load_workbook(_as_workbook_source(source), read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        return {}
    worksheet = workbook[sheet_name]
    mappings: dict[str, str | None] = {}
    for row in worksheet.iter_rows(min_row=4, values_only=True):
        code = row[0] if len(row) > 0 else None
        name = row[1] if len(row) > 1 else None
        if name is None or code in (None, "", 0, "0"):
            continue
        mappings[" ".join(str(name).casefold().split())] = str(code)
    return mappings


def _find_trial_balance_sheet(workbook: openpyxl.Workbook) -> openpyxl.worksheet.worksheet.Worksheet:
    for worksheet in workbook.worksheets:
        title = worksheet.title.casefold()
        if "tb" in title or "trial" in title or "tally" in title:
            return worksheet
    return workbook.worksheets[0]


def _as_workbook_source(source: str | Path | bytes | BinaryIO) -> str | Path | BinaryIO:
    return BytesIO(source) if isinstance(source, bytes) else source


def _cell_depth(cell: object) -> int:
    alignment = getattr(cell, "alignment", None)
    indent = getattr(alignment, "indent", 0) or 0
    return int(indent)